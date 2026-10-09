#!/usr/bin/env python3
"""Import the društva lead list (xlsx) into the `drustvo` table.

  DRUSTVA_EMAIL=… DRUSTVA_PASSWORD=… python3 scripts/drustva-import.py ~/Downloads/drustva_seznam_v1.xlsx [--dry]

Signs in as a staff user (RLS applies), reads the main sheet "Društva" and the
"Umbrella federations" sheet, and inserts only NEW e-mail addresses — re-running
never overwrites edits made in Colnix. URL and anon key come from .env.local.
"""
import json, os, re, sys, urllib.request, urllib.error
import openpyxl

TIER = {"Focus": "focus", "50/50": "fifty_fifty", "Not chosen": "not_chosen"}
CHECK = {
    "Confirmed (source page)": "confirmed",
    "Confirmed (source page + own website)": "confirmed_own_site",
    "Updated to email from own website": "updated",
    "Not sure": "not_sure",
    "Probably wrong": "probably_wrong",
}


def env_local():
    out = {}
    path = os.path.join(os.path.dirname(__file__), "..", ".env.local")
    for line in open(path, encoding="utf-8"):
        m = re.match(r"\s*([A-Z0-9_]+)\s*=\s*(.*)\s*$", line)
        if m:
            out[m.group(1)] = m.group(2).strip().strip('"').strip("'")
    return out


def clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def lower_email(v):
    s = clean(v)
    return s.lower() if s else None


def to_int(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


def read_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, seen = [], set()

    ws = wb["Društva"]
    data = list(ws.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(data[0])}
    dist_col = next(k for k in idx if k and k.startswith("Distance from"))
    for r in data[1:]:
        email = lower_email(r[idx["Email"]])
        if not email or email in seen:
            continue
        seen.add(email)
        rows.append({
            "name": clean(r[idx["Name"]]),
            "type": clean(r[idx["Type"]]),
            "town": clean(r[idx["Town"]]),
            "region": clean(r[idx["Region"]]),
            "tier": TIER.get(clean(r[idx["Priority tier"]]), "not_chosen"),
            "email": email,
            "email_alt": lower_email(r[idx["Old / alternative email"]]),
            "website": clean(r[idx["Website / Facebook"]]),
            "email_check": CHECK.get(clean(r[idx["Email check"]])),
            "distance_km": to_int(r[idx[dist_col]]),
            "distance_band": clean(r[idx["Distance band"]]),
            "activity_level": (clean(r[idx["Activity level"]]) or "unknown").lower(),
            "organizes_trips": (clean(r[idx["Organizes trips?"]]) or "unknown").lower(),
            "activity_note": clean(r[idx["Activity note (what we saw)"]]),
            "source_url": clean(r[idx["Evidence link"]]),
            "email_source_url": clean(r[idx["Email source page"]]),
            "source_round": clean(r[idx["Found in"]]),
            "notes": clean(r[idx["Notes"]]),
        })

    if "Umbrella federations" in wb.sheetnames:
        ws = wb["Umbrella federations"]
        data = list(ws.iter_rows(values_only=True))
        idx = {h: i for i, h in enumerate(data[0])}
        for r in data[1:]:
            email = lower_email(r[idx["Email"]])
            if not email or email in seen or not clean(r[idx["Name"]]):
                continue
            seen.add(email)
            rows.append({
                "name": clean(r[idx["Name"]]),
                "type": "Krovna zveza",
                "tier": "not_chosen",
                "email": email,
                "source_url": clean(r[idx["Source"]]),
                "notes": f"Pokriva: {clean(r[idx['Covers']]) or '—'}",
            })
    # PostgREST wants identical keys in every row of a batch
    keys = sorted({k for r in rows for k in r})
    return [{k: r.get(k) for k in keys} for r in rows]


def request(url, key, token=None, method="GET", body=None, headers=None):
    h = {"apikey": key, "Content-Type": "application/json"}
    if token:
        h["Authorization"] = f"Bearer {token}"
    h.update(headers or {})
    req = urllib.request.Request(url, data=json.dumps(body).encode() if body is not None else None, headers=h, method=method)
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.status, r.read().decode()
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode()


if __name__ == "__main__":
    path = sys.argv[1]
    dry = "--dry" in sys.argv
    rows = read_rows(path)
    print(f"{len(rows)} rows read")
    if dry:
        print(json.dumps(rows[0], ensure_ascii=False, indent=1))
        sys.exit(0)

    cfg = env_local()
    base, anon = cfg["NEXT_PUBLIC_SUPABASE_URL"], cfg["NEXT_PUBLIC_SUPABASE_ANON_KEY"]
    status, text = request(f"{base}/auth/v1/token?grant_type=password", anon, method="POST",
                           body={"email": os.environ["DRUSTVA_EMAIL"], "password": os.environ["DRUSTVA_PASSWORD"]})
    if status != 200:
        sys.exit(f"login failed: {status}")
    token = json.loads(text)["access_token"]

    added = 0
    for i in range(0, len(rows), 200):
        chunk = rows[i:i + 200]
        status, text = request(f"{base}/rest/v1/drustvo?on_conflict=email", anon, token, "POST", chunk,
                               {"Prefer": "resolution=ignore-duplicates,return=minimal"})
        if status not in (200, 201, 204):
            sys.exit(f"insert failed at {i}: {status} {text[:300]}")
        added += len(chunk)
        print(f"sent {added}/{len(rows)}", flush=True)
    print("done")
