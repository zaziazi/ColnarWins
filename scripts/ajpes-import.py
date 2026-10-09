#!/usr/bin/env python3
"""
Finds hospitality businesses (SKD 55.x accommodation, 56.x food & drink) in an
AJPES / Bisnode company-register export and works out which ones are in a
region. The export has NO town or post code, only street + number, so the
location is derived by geocoding inside the region's bounding box and then
cross-checked.

  python3 scripts/ajpes-import.py extract  <register.xlsx>  hosp.json
  python3 scripts/ajpes-import.py locate   hosp.json  located.json \
          --bbox 45.55,14.7,46.15,15.55 --customer-vats dolenjska-vat.json
  python3 scripts/ajpes-import.py sql      located.json  ./ajpes-sql --region dolenjska [--unverified]

`--customer-vats` is {"<8-digit VAT>": "<post code>"} for our customers in the
region; a business whose VAT matches one of them is placed with certainty.

Nominatim policy: identifying User-Agent, at most 1 request per second.
"""
import json, re, sys, time, urllib.parse, urllib.request

UA = "colnix-sales-map/1.0 (colnar.aljaz.ac@gmail.com)"

# Landline prefixes that belong to the region (07 = Dolenjska/Bela krajina/Posavje;
# the 01 numbers are Grosuplje, Ivančna Gorica, Kočevje, Ribnica).
REGION_PHONE = re.compile(r"^(07|01(76|78|83|89)|0?1\s?(76|78|83|89))")
REGION_TOWNS = [
    "novo mesto", "trebnje", "metlika", "črnomelj", "kočevje", "ribnica", "grosuplje",
    "ivančna gorica", "šentjernej", "žužemberk", "dolenjske toplice", "otočec", "šmarješke toplice",
    "mirna peč", "mokronog", "straža", "semič", "krško", "sevnica", "čatež", "kostanjevica",
    "škocjan", "bela krajina", "dolenjska", "dobrnič", "stopiče", "brusnice", "podbočje",
]


def num(v):
    if v is None:
        return None
    s = str(v).replace(".", "").replace(",", ".").strip()
    try:
        return float(s)
    except ValueError:
        return None


def extract(xlsx, out):
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    rows_out = []
    for name in ["d.d+d.o.o.", "Podružnice", "Zadruge, društva,..", "s.p.", "Agencije"]:
        if name not in wb.sheetnames:
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        h = next(i for i, r in enumerate(rows[:60]) if r and "Firma" in [str(c) for c in r if c is not None])
        col = {str(c).strip(): i for i, c in enumerate(rows[h]) if c is not None}
        for r in rows[h + 1 :]:
            if not r[col["Firma"]]:
                continue
            try:
                skd = float(r[col["Dejavnost SKD"]])
            except (TypeError, ValueError):
                continue
            if int(skd) not in (55, 56):
                continue
            vat = r[col["Davčna številka"]]
            rows_out.append(
                dict(
                    sheet=name,
                    firma=str(r[col["Firma"]]).strip(),
                    naslov=(str(r[col["Naslov / Sedež"]]).strip() if r[col["Naslov / Sedež"]] else None),
                    maticna=str(int(r[col["Matična številka"]])) if r[col["Matična številka"]] else None,
                    vat=str(int(vat)).zfill(8) if vat else None,
                    tel=(str(r[col["Telefon"]]).strip() if r[col["Telefon"]] else None),
                    email=(str(r[col["Email"]]).strip().lower() if r[col["Email"]] else None),
                    zastopnik=r[col["Zastopnik"]],
                    skd=skd,
                    prihodki=num(r[col["Skupni prihodki"]]),
                    zaposleni=num(r[col["Število zaposlenih"]]),
                )
            )
    json.dump(rows_out, open(out, "w"), ensure_ascii=False)
    print(f"{len(rows_out)} hospitality businesses -> {out}")


_last = 0.0


def nominatim(params):
    global _last
    wait = 1.1 - (time.time() - _last)
    if wait > 0:
        time.sleep(wait)
    url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                _last = time.time()
                return json.load(r)
        except Exception as e:  # noqa: BLE001
            _last = time.time()
            time.sleep(3 * (attempt + 1))
            err = e
    raise err


def km(a, b):
    from math import asin, cos, radians, sin, sqrt

    dlat, dlng = radians(b[0] - a[0]), radians(b[1] - a[1])
    h = sin(dlat / 2) ** 2 + cos(radians(a[0])) * cos(radians(b[0])) * sin(dlng / 2) ** 2
    return 2 * 6371 * asin(sqrt(h))


def locate(inp, out, bbox, vats):
    s, w, n, e = [float(x) for x in bbox.split(",")]
    viewbox = f"{w},{n},{e},{s}"
    items = json.load(open(inp))
    vat_post = json.load(open(vats)) if vats else {}
    done = {}
    try:
        done = {d["maticna"]: d for d in json.load(open(out))}
    except Exception:  # noqa: BLE001
        pass
    results = list(done.values())
    for i, it in enumerate(items):
        if it["maticna"] in done:
            continue
        addr = it["naslov"] or ""
        res = dict(it, found=False, strong=False, evidence=[])
        if addr:
            q = addr + ", Slovenija"
            hit = nominatim(dict(q=q, format="jsonv2", limit=1, countrycodes="si", viewbox=viewbox,
                                 bounded=1, addressdetails=1))
            if hit:
                h = hit[0]
                res.update(found=True, lat=float(h["lat"]), lng=float(h["lon"]), rank=h.get("place_rank"),
                           display=h.get("display_name"), addr=h.get("address", {}))
                free = nominatim(dict(q=q, format="jsonv2", limit=1, countrycodes="si"))
                if free and km((res["lat"], res["lng"]), (float(free[0]["lat"]), float(free[0]["lon"]))) < 1.0:
                    res["evidence"].append("unique")  # the best match anywhere in Slovenia is this one
        tel = re.sub(r"[^\d]", "", it["tel"] or "")
        if tel.startswith("386"):
            tel = "0" + tel[3:]
        if tel and REGION_PHONE.match(tel):
            res["evidence"].append("phone")
        if it["vat"] and it["vat"] in vat_post:
            res["evidence"].append("customer-vat")
        text = f"{it['firma']} {addr}".lower()
        if any(t in text for t in REGION_TOWNS):
            res["evidence"].append("town-name")
        res["strong"] = bool(res["found"] and res["evidence"])
        if not res["found"] and "customer-vat" in res["evidence"]:
            res["strong"] = True  # in the region by VAT even if the street didn't geocode
        results.append(res)
        if (i + 1) % 25 == 0:
            json.dump(results, open(out, "w"), ensure_ascii=False)
            print(f"{i + 1}/{len(items)}", flush=True)
    json.dump(results, open(out, "w"), ensure_ascii=False)
    print("done", len(results), "found", sum(r["found"] for r in results), "strong", sum(r["strong"] for r in results))


KIND_BY_SKD = {
    56.111: "restaurant", 56.112: "restaurant", 56.113: "fast_food", 56.12: "fast_food",
    56.21: "catering", 56.22: "catering", 56.3: "bar", 55.1: "hotel",
    55.201: "guest_house", 55.202: "guest_house", 55.203: "guest_house", 55.204: "guest_house",
    55.209: "guest_house", 55.3: "camping",
}

LEGAL_FORM = re.compile(r"(?:,\s*|\s+)(d\.?\s?o\.?\s?o\.?|d\.?\s?d\.?|d\.?\s?n\.?\s?o\.?|k\.?\s?d\.?|z\.?\s?o\.?\s?o\.?)(?![a-zčšž]).*$", re.I)
SOLE_TRADER = re.compile(r"\s+s\.?\s?p\.?\s*$", re.I)


def display_name(firma):
    """Registered name -> a name people would recognise on a map.

    'TERME KRKA, d.o.o., Novo mesto'          -> 'Terme Krka'
    'GOSTILNA PRI JAPU, KEPA JOŠKO S.P.'      -> 'Gostilna Pri Japu'   (trade name before the owner)
    'BLAŽ MRHAR S.P.'                         -> 'Blaž Mrhar s.p.'      (only the owner's name is known)
    """
    n = firma.strip().strip('"')
    if SOLE_TRADER.search(n):
        base = SOLE_TRADER.sub("", n).strip(" ,")
        if "," in base:
            n = base.split(",")[0].strip()
            suffix = ""
        else:
            n = base
            suffix = " s.p."
    else:
        n = LEGAL_FORM.sub("", n).strip(" ,")
        suffix = ""
    if n.isupper() and len(n) > 3:
        n = n.title()
    return n + suffix


def sql_literal(v):
    return "null" if v is None else "'" + str(v).replace("'", "''") + "'"


def build_sql(inp, outdir, region, unverified=False):
    import os

    os.makedirs(outdir, exist_ok=True)
    rows = [
        r for r in json.load(open(inp))
        if r.get("found") and (r.get("unverified") if unverified else r.get("strong"))
    ]
    status = "unverified" if unverified else "ok"
    cands = []
    for r in rows:
        addr = r.get("addr", {})
        road = addr.get("road") or addr.get("pedestrian") or ""
        street = " ".join(x for x in [road, addr.get("house_number")] if x) or None
        city = addr.get("city") or addr.get("town") or addr.get("village") or addr.get("municipality")
        tel = (r.get("tel") or "").split(";")[0].strip() or None
        cands.append(
            dict(
                ajpes_id=r["maticna"], vat_id=r.get("vat"), legal_name=r["firma"], name=display_name(r["firma"]),
                kind=KIND_BY_SKD.get(round(r["skd"], 3), KIND_BY_SKD.get(round(r["skd"], 2), "restaurant")),
                lat=r["lat"], lng=r["lng"], address=street or r.get("naslov"), city=city,
                post_code=addr.get("postcode"), phone=tel, email=r.get("email"),
                representative=r.get("zastopnik"), revenue_eur=r.get("prihodki"),
                employees=int(r["zaposleni"]) if r.get("zaposleni") is not None else None,
                skd_code=str(r["skd"]),
            )
        )
    cols = ("ajpes_id text, vat_id text, legal_name text, name text, kind text, lat double precision, lng double precision, "
            "address text, city text, post_code text, phone text, email text, representative text, revenue_eur numeric, "
            "employees int, skd_code text")
    for n, i in enumerate(range(0, len(cands), 60), 1):
        payload = json.dumps(cands[i : i + 60], ensure_ascii=False).replace("'", "''")
        sql = f"""with cand as (
  select * from jsonb_to_recordset('{payload}'::jsonb) as x({cols})
), matched as (
  -- already on the map (from OpenStreetMap)? same place = close by and a similar name
  select distinct on (c.ajpes_id) c.ajpes_id, v.id as venue_id
    from cand c
    join public.venue v on v.ajpes_id is null and v.region = '{region}'
     and abs(v.lat - c.lat) < 0.0015 and abs(v.lng - c.lng) < 0.002
     and public.geo_distance_m(v.lat, v.lng, c.lat, c.lng) <= 120
     and greatest(
           similarity(public.norm_name(v.name), public.norm_name(c.name)),
           similarity(public.norm_name(v.name), public.norm_name(c.legal_name)),
           word_similarity(public.norm_name(v.name), public.norm_name(c.legal_name))) >= 0.3
   order by c.ajpes_id, public.geo_distance_m(v.lat, v.lng, c.lat, c.lng)
), upd as (
  update public.venue v set
    ajpes_id = c.ajpes_id, vat_id = c.vat_id, legal_name = c.legal_name, representative = c.representative,
    revenue_eur = c.revenue_eur, employees = c.employees, skd_code = c.skd_code,
    phone = coalesce(v.phone, c.phone), email = coalesce(v.email, c.email)
   from matched m join cand c on c.ajpes_id = m.ajpes_id
  where v.id = m.venue_id
  returning v.id
)
insert into public.venue (source, ajpes_id, vat_id, legal_name, name, kind, lat, lng, address, city, post_code, phone, email,
                          representative, revenue_eur, employees, skd_code, region, location_status)
select 'ajpes', c.ajpes_id, c.vat_id, c.legal_name, c.name, c.kind, c.lat, c.lng, c.address, c.city, c.post_code, c.phone, c.email,
       c.representative, c.revenue_eur, c.employees, c.skd_code, '{region}', '{status}'
  from cand c
 where c.ajpes_id not in (select ajpes_id from matched)
on conflict (ajpes_id) where ajpes_id is not null do update set
  vat_id = excluded.vat_id, legal_name = excluded.legal_name, representative = excluded.representative,
  revenue_eur = excluded.revenue_eur, employees = excluded.employees, skd_code = excluded.skd_code;
"""
        open(f"{outdir}/ajpes-{n}.sql", "w", encoding="utf-8").write(sql)
    print(f"{len(cands)} businesses -> {outdir}")


if __name__ == "__main__":
    cmd = sys.argv[1]
    if cmd == "extract":
        extract(sys.argv[2], sys.argv[3])
    elif cmd == "sql":
        a = sys.argv
        build_sql(a[2], a[3], a[a.index("--region") + 1] if "--region" in a else "dolenjska", "--unverified" in a)
    elif cmd == "locate":
        a = sys.argv
        opt = lambda k: a[a.index(k) + 1] if k in a else None  # noqa: E731
        locate(a[2], a[3], opt("--bbox") or "45.55,14.7,46.15,15.55", opt("--customer-vats"))
